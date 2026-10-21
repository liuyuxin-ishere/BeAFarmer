#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from openpyxl import Workbook
from pathlib import Path

DEFINES = Path(__file__).resolve().parents[1] / "Shared" / "config" / "Defines"
DEFINES.mkdir(parents=True, exist_ok=True)

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append(
    [
        "##var",
        "full_name",
        "value_type",
        "read_schema_from_file",
        "input",
        "index",
        "mode",
        "group",
        "comment",
        "tags",
        "output",
    ]
)
ws.append(["##", "全名", "记录类名", "从excel读取定义", "文件列表", "表id字段", "模式", "分组", "注释", None, "输出文件名"])
ws.append(["##", None, None, "true=从数据表头读", None, "联合主键 key1+key2", "one|map|list", None, None, None, None])
rows = [
    ["TbGlobal", "Global", True, "global.xlsx", None, "one", None, "全局", None, None],
    ["TbItem", "Item", True, "item.xlsx", "id", "map", None, "物品", None, None],
    ["TbCrop", "Crop", True, "crop.xlsx", "id", "map", None, "作物", None, None],
    ["TbGatherNode", "GatherNode", True, "gathernode.xlsx", "id", "map", None, "采集点", None, None],
    ["TbTool", "Tool", True, "tool.xlsx", "toolId+level", "map", None, "工具", None, None],
    ["TbFacility", "Facility", True, "facility.xlsx", "id", "map", None, "设施", None, None],
    ["TbFacilityLevel", "FacilityLevel", True, "facilitylevel.xlsx", "level", "map", None, "设施等级", None, None],
    ["TbRecipe", "Recipe", True, "recipe.xlsx", "id", "map", None, "配方", None, None],
    ["TbLevel", "Level", True, "level.xlsx", "level", "map", None, "等级", None, None],
    ["TbPlot", "Plot", True, "plot.xlsx", "stage", "map", None, "地块", None, None],
    ["TbOrderRule", "OrderRule", True, "orderrule.xlsx", "key", "map", None, "订单规则", None, None],
    ["TbFurniture", "Furniture", True, "furniture.xlsx", "id", "map", None, "家具", None, None],
    ["TbAvatarItem", "AvatarItem", True, "avataritem.xlsx", "id", "map", None, "外观", None, None],
    ["TbGachaRule", "GachaRule", True, "gacharule.xlsx", "key", "map", None, "抽卡规则", None, None],
    ["TbGachaPool", "GachaPool", True, "gachapool.xlsx", "id", "map", None, "卡池", None, None],
    ["TbGachaPoolItem", "GachaPoolItem", True, "gachapoolitem.xlsx", None, "list", None, "卡池物品", None, None],
    ["TbActivity", "Activity", True, "activity.xlsx", "id", "map", None, "活动", None, None],
]
for r in rows:
    ws.append([None, *r])
wb.save(DEFINES / "__tables__.xlsx")

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append(
    [
        "##var",
        "full_name",
        "parent",
        "valueType",
        "alias",
        "sep",
        "comment",
        "tags",
        "group",
        "*fields",
        None,
        None,
        None,
        None,
        None,
        None,
    ]
)
ws.append(
    [
        "##var",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "name",
        "alias",
        "type",
        "group",
        "comment",
        "tags",
        "variants",
    ]
)
ws.append(
    [
        "##",
        "全名",
        "父类",
        "是否值类型",
        "别名",
        "分割符",
        "注释",
        None,
        None,
        "字段名",
        "别名",
        "类型",
        "分组",
        "注释",
        None,
        None,
    ]
)
ws.append([None, "ItemCount", None, None, None, ":", "物品数量", None, None, "id", None, "int", None, "物品ID", None, None])
ws.append([None, None, None, None, None, None, None, None, None, "count", None, "int", None, "数量", None, None])
wb.save(DEFINES / "__beans__.xlsx")

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append(["##var", "full_name", "flags", "unique", "group", "comment", "tags", "*items", None, None, None, None])
ws.append(["##var", None, None, None, None, None, None, "name", "alias", "value", "comment", "tags"])
ws.append(["##", "全名", "flags", "unique", None, "注释", None, "枚举名", "别名", "值", "注释", None])


def add_enum(name, items):
    for i, (n, v) in enumerate(items):
        if i == 0:
            ws.append([None, name, False, True, None, None, None, n, None, v, None, None])
        else:
            ws.append([None, None, None, None, None, None, None, n, None, v, None, None])


add_enum(
    "ItemType",
    [
        ("Seed", 0),
        ("Crop", 1),
        ("Craft", 2),
        ("Gather", 3),
        ("Tool", 4),
        ("Consumable", 5),
        ("Furniture", 6),
        ("Currency", 7),
    ],
)
add_enum("AvatarPart", [("Hair", 0), ("Top", 1), ("Bottom", 2), ("Shoes", 3), ("Hat", 4), ("Acc", 5), ("Dress", 6)])
add_enum("Rarity", [("N", 0), ("R", 1), ("SR", 2), ("SSR", 3)])
add_enum("AvatarSource", [("Initial", 0), ("Shop", 1), ("Gacha", 2)])
add_enum("PoolType", [("Permanent", 0), ("Limited", 1)])
wb.save(DEFINES / "__enums__.xlsx")
print("defines ok", DEFINES)
