#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PoC #7：从 Excel 导表 → JSON 数据文件 + C# 行类型 / 表加载代码。

约定（与 xlsx「说明」页一致）：
  第 1 行字段名，第 2 行类型，第 3 行注释，第 4 行起数据。
  字段名以 # 开头或类型为 calc 的列跳过；空字段名跳过。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "Assets" / "ReferenceDoc" / "云屿小筑_配置表与开发进度.xlsx"
JSON_OUT = ROOT / "Assets" / "StreamingAssets" / "Config"
CS_OUT = ROOT / "Assets" / "Shared" / "Config" / "Generated"
SERVER_JSON_OUT = ROOT / "Server" / "Config"

SKIP_SHEETS = {"说明", "校验", "开发进度", "OrderCalc", "GachaSim"}

CS_TYPE = {
    "int": "int",
    "long": "long",
    "float": "float",
    "double": "double",
    "bool": "bool",
    "string": "string",
}


def to_pascal(name: str) -> str:
    parts = re.split(r"[_\s]+", name.strip())
    return "".join(p[:1].upper() + p[1:] if p else "" for p in parts)


def should_export_col(field, typ) -> bool:
    if field is None:
        return False
    field = str(field).strip()
    if not field or field.startswith("#"):
        return False
    if typ is None:
        return False
    typ = str(typ).strip().lower()
    if typ in ("", "calc"):
        return False
    return typ in CS_TYPE


def convert_value(typ: str, raw):
    if raw is None:
        if typ == "string":
            return ""
        if typ == "bool":
            return False
        if typ in ("int", "long"):
            return 0
        if typ in ("float", "double"):
            return 0.0
        return None
    if typ == "string":
        return str(raw)
    if typ == "bool":
        if isinstance(raw, bool):
            return raw
        s = str(raw).strip().lower()
        return s in ("1", "true", "yes", "y")
    if typ == "int":
        return int(round(float(raw)))
    if typ == "long":
        return int(round(float(raw)))
    if typ == "float":
        return float(raw)
    if typ == "double":
        return float(raw)
    return raw


def parse_sheet(ws):
    max_col = ws.max_column or 0
    max_row = ws.max_row or 0
    fields = [ws.cell(1, c).value for c in range(1, max_col + 1)]
    types = [ws.cell(2, c).value for c in range(1, max_col + 1)]
    comments = [ws.cell(3, c).value for c in range(1, max_col + 1)]

    cols = []
    for i, (f, t, cmt) in enumerate(zip(fields, types, comments)):
        if not should_export_col(f, t):
            continue
        fname = str(f).strip()
        tname = str(t).strip().lower()
        cols.append(
            {
                "index": i + 1,
                "name": fname,
                "pascal": to_pascal(fname),
                "type": tname,
                "cs": CS_TYPE[tname],
                "comment": "" if cmt is None else str(cmt),
            }
        )

    if not cols:
        return None

    rows = []
    for r in range(4, max_row + 1):
        # skip fully empty rows (all export cols empty)
        values = [ws.cell(r, col["index"]).value for col in cols]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        obj = {}
        for col, raw in zip(cols, values):
            obj[col["name"]] = convert_value(col["type"], raw)
        rows.append(obj)

    # uniqueness for id-like keys
    for key in ("id", "key", "toolId", "poolId", "activityId", "level", "stage"):
        if any(c["name"] == key for c in cols):
            seen = set()
            # composite uniqueness for Tool (toolId+level) and GachaPoolItem
            if key == "toolId" and any(c["name"] == "level" for c in cols):
                for row in rows:
                    sig = (row.get("toolId"), row.get("level"))
                    if sig in seen:
                        raise ValueError(f"{ws.title}: duplicate toolId+level {sig}")
                    seen.add(sig)
            elif key == "poolId" and any(c["name"] == "itemId" for c in cols):
                for row in rows:
                    sig = (row.get("poolId"), row.get("itemId"), row.get("rarity"))
                    if sig in seen:
                        raise ValueError(f"{ws.title}: duplicate pool item {sig}")
                    seen.add(sig)
            elif key in ("id", "key", "activityId") or (
                key == "poolId" and not any(c["name"] == "itemId" for c in cols)
            ) or key in ("level", "stage") and ws.title in ("Level", "Plot", "FacilityLevel"):
                for row in rows:
                    v = row.get(key)
                    if v in seen:
                        raise ValueError(f"{ws.title}: duplicate {key}={v}")
                    seen.add(v)
            break

    return {"name": ws.title, "cols": cols, "rows": rows}


def emit_row_class(table) -> str:
    name = table["name"]
    cls = f"{name}Cfg"
    lines = [
        "    /// <summary>Generated row for table " + name + ".</summary>",
        f"    public sealed class {cls}",
        "    {",
    ]
    for col in table["cols"]:
        if col["comment"]:
            lines.append(f'        /// <summary>{col["comment"]}</summary>')
        lines.append(f'        public {col["cs"]} {col["pascal"]} {{ get; set; }}')
    lines.append("    }")
    lines.append("")
    return "\n".join(lines)


def pick_unique_key(table):
    """仅当候选字段在全表取值唯一时才作为 Get 主键。"""
    cols = table["cols"]
    rows = table["rows"]
    name_to_col = {c["name"]: c for c in cols}
    for cand in ("id", "key", "activityId", "poolId", "level", "stage", "toolId"):
        if cand not in name_to_col:
            continue
        values = [r.get(cand) for r in rows]
        if len(values) == len(set(values)):
            return name_to_col[cand]
    return None


def emit_table_class(table) -> str:
    name = table["name"]
    cls = f"{name}Cfg"
    tbl = f"Tb{name}"
    key_col = pick_unique_key(table)
    has_key = key_col is not None

    lines = [
        f"    /// <summary>Generated table {name} ({len(table['rows'])} rows).</summary>",
        f"    public sealed class {tbl}",
        "    {",
        f"        readonly System.Collections.Generic.List<{cls}> _list = new System.Collections.Generic.List<{cls}>();",
    ]
    if has_key:
        lines.append(
            f"        readonly System.Collections.Generic.Dictionary<{key_col['cs']}, {cls}> _byKey "
            f"= new System.Collections.Generic.Dictionary<{key_col['cs']}, {cls}>();"
        )
    lines.append("")
    lines.append(f"        public System.Collections.Generic.IReadOnlyList<{cls}> All => _list;")
    lines.append(f"        public int Count => _list.Count;")
    lines.append("")

    if has_key:
        lines.append(f"        public {cls} Get({key_col['cs']} key)")
        lines.append("        {")
        lines.append("            if (_byKey.TryGetValue(key, out var row)) return row;")
        lines.append(f'            throw new System.Collections.Generic.KeyNotFoundException("{name} missing key=" + key);')
        lines.append("        }")
        lines.append("")
        lines.append(f"        public bool TryGet({key_col['cs']} key, out {cls} row) => _byKey.TryGetValue(key, out row);")
        lines.append("")

    if name == "Global" or name == "OrderRule" or name == "GachaRule":
        lines.append("        public string GetString(string key, string defaultValue = \"\")")
        lines.append("        {")
        lines.append("            if (TryGet(key, out var row) && row.Value != null) return row.Value;")
        lines.append("            return defaultValue;")
        lines.append("        }")
        lines.append("")
        lines.append("        public int GetInt(string key, int defaultValue = 0)")
        lines.append("        {")
        lines.append("            var s = GetString(key, null);")
        lines.append("            if (s == null) return defaultValue;")
        lines.append("            return int.TryParse(s, System.Globalization.NumberStyles.Integer, System.Globalization.CultureInfo.InvariantCulture, out var v)")
        lines.append("                ? v")
        lines.append("                : (int)System.Math.Round(double.Parse(s, System.Globalization.CultureInfo.InvariantCulture));")
        lines.append("        }")
        lines.append("")
        lines.append("        public float GetFloat(string key, float defaultValue = 0f)")
        lines.append("        {")
        lines.append("            var s = GetString(key, null);")
        lines.append("            if (s == null) return defaultValue;")
        lines.append("            return float.Parse(s, System.Globalization.CultureInfo.InvariantCulture);")
        lines.append("        }")
        lines.append("")

    lines.append(f"        public static {tbl} LoadFromJson(string json)")
    lines.append("        {")
    lines.append(f"            var table = new {tbl}();")
    lines.append("            var arr = CloudIsle.Config.SimpleJson.ParseArray(json);")
    lines.append("            for (int i = 0; i < arr.Count; i++)")
    lines.append("            {")
    lines.append("                var obj = arr[i];")
    lines.append(f"                var row = new {cls}();")
    for col in table["cols"]:
        lines.append(emit_assign(col))
    lines.append("                table._list.Add(row);")
    if has_key:
        lines.append(f"                table._byKey[row.{key_col['pascal']}] = row;")
    lines.append("            }")
    lines.append("            return table;")
    lines.append("        }")
    lines.append("    }")
    lines.append("")
    return "\n".join(lines)


def emit_assign(col) -> str:
    n = col["name"]
    p = col["pascal"]
    t = col["cs"]
    if t == "string":
        return f'                row.{p} = obj.GetString("{n}");'
    if t == "int":
        return f'                row.{p} = obj.GetInt("{n}");'
    if t == "long":
        return f'                row.{p} = obj.GetLong("{n}");'
    if t == "float":
        return f'                row.{p} = obj.GetFloat("{n}");'
    if t == "double":
        return f'                row.{p} = obj.GetDouble("{n}");'
    if t == "bool":
        return f'                row.{p} = obj.GetBool("{n}");'
    return f'                // unsupported {n}'


def emit_tables_root(tables) -> str:
    lines = [
        "    public sealed class ConfigTables",
        "    {",
    ]
    for t in tables:
        lines.append(f"        public Tb{t['name']} {t['name']} {{ get; private set; }}")
    lines.append("")
    lines.append("        public static ConfigTables LoadFromDirectory(string directory)")
    lines.append("        {")
    lines.append("            if (string.IsNullOrEmpty(directory) || !System.IO.Directory.Exists(directory))")
    lines.append('                throw new System.IO.DirectoryNotFoundException("config dir: " + directory);')
    lines.append("            var tables = new ConfigTables();")
    for t in tables:
        fname = t["name"].lower() + ".json"
        lines.append(f'            tables.{t["name"]} = Tb{t["name"]}.LoadFromJson(System.IO.File.ReadAllText(System.IO.Path.Combine(directory, "{fname}")));')
    lines.append("            return tables;")
    lines.append("        }")
    lines.append("    }")
    lines.append("")
    return "\n".join(lines)


def write_generated_cs(tables):
    CS_OUT.mkdir(parents=True, exist_ok=True)
    parts = [
        "// <auto-generated>",
        "// PoC #7 config tables — DO NOT EDIT. Run Tools/gen_config.py",
        "// </auto-generated>",
        "",
        "namespace CloudIsle.Config",
        "{",
        "",
    ]
    for t in tables:
        parts.append(emit_row_class(t))
        parts.append(emit_table_class(t))
    parts.append(emit_tables_root(tables))
    parts.append("}")
    parts.append("")
    path = CS_OUT / "ConfigTables.g.cs"
    path.write_text("\n".join(parts), encoding="utf-8")
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--set-global", action="append", default=[], help="key=value override before export (PoC verify)")
    args = ap.parse_args()

    xlsx = args.xlsx
    if not xlsx.is_file():
        print(f"ERROR: missing xlsx: {xlsx}", file=sys.stderr)
        return 1

    # data_only for formula sheets we skip; values for config sheets are mostly literals
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    # optional in-memory overrides on Global (does not write back xlsx)
    overrides = {}
    for item in args.set_global:
        if "=" not in item:
            print(f"ERROR: bad --set-global {item}", file=sys.stderr)
            return 1
        k, v = item.split("=", 1)
        overrides[k] = v

    tables = []
    for name in wb.sheetnames:
        if name in SKIP_SHEETS:
            continue
        parsed = parse_sheet(wb[name])
        if parsed is None:
            print(f"SKIP empty/non-exportable: {name}")
            continue
        if name == "Global" and overrides:
            for row in parsed["rows"]:
                if row["key"] in overrides:
                    row["value"] = overrides[row["key"]]
                    print(f"OVERRIDE Global {row['key']}={row['value']}")
        tables.append(parsed)
        print(f"OK {name}: {len(parsed['rows'])} rows, {len(parsed['cols'])} cols")

    JSON_OUT.mkdir(parents=True, exist_ok=True)
    SERVER_JSON_OUT.mkdir(parents=True, exist_ok=True)
    for t in tables:
        payload = t["rows"]
        text = json.dumps(payload, ensure_ascii=False, indent=2) + "\n"
        fname = t["name"].lower() + ".json"
        (JSON_OUT / fname).write_text(text, encoding="utf-8")
        (SERVER_JSON_OUT / fname).write_text(text, encoding="utf-8")

    cs_path = write_generated_cs(tables)
    manifest = {
        "source": str(xlsx.relative_to(ROOT)).replace("\\", "/"),
        "tables": [{ "name": t["name"], "rows": len(t["rows"]), "cols": [c["name"] for c in t["cols"]] } for t in tables],
    }
    (JSON_OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (SERVER_JSON_OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"JSON -> {JSON_OUT}")
    print(f"JSON -> {SERVER_JSON_OUT}")
    print(f"C#  -> {cs_path}")
    print(f"TABLES {len(tables)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
